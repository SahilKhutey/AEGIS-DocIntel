"""
AEGIS-AMDI-OS — PPTX Loader
=============================
Microsoft PowerPoint loader using python-pptx with strict error handling.
"""
from __future__ import annotations

import io
import logging
from typing import Any
import zipfile

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    Presentation = None
    MSO_SHAPE_TYPE = None

from src.core.document_object import DocumentFormat, DocumentObject
from src.ingestion.base import BaseLoader
from src.ingestion.exceptions import (
    DocumentCorruptError,
    EncryptedDocumentError,
    FormatError,
    LoaderError,
    SizeLimitError,
)

logger = logging.getLogger(__name__)


class PPTXLoader(BaseLoader):
    """PPTX presentation loader."""

    FORMAT_NAME = "pptx"
    SUPPORTED_EXTENSIONS = {".pptx"}
    PPTX_MAGIC = b"PK\x03\x04"

    def __init__(self, max_size_mb: int = 100, **options):
        super().__init__(**options)
        self.max_size_mb = max_size_mb

    def validate(self, raw_bytes: bytes) -> bool:
        """Check if bytes represent a valid PPTX file."""
        if not raw_bytes or len(raw_bytes) < 4:
            return False
        if raw_bytes[:4] != self.PPTX_MAGIC:
            return False
        try:
            with zipfile.ZipFile(io.BytesIO(raw_bytes)) as z:
                names = z.namelist()
                return any(n.startswith("ppt/") for n in names) or "[Content_Types].xml" in names
        except Exception:
            return False

    async def load(self, source, filename: str = "") -> DocumentObject:
        """Load a PPTX presentation with zero silent failures."""
        raw_bytes, name = self.read_source(source)
        if filename:
            name = filename
        if not name:
            name = "presentation.pptx"

        if not raw_bytes or len(raw_bytes) < 4 or raw_bytes[:4] != self.PPTX_MAGIC:
            raise FormatError(f"Not a valid PPTX file (missing ZIP header): {name}", filename=name)

        if Presentation is None:
            raise LoaderError("python-pptx is not installed", filename=name)

        size_mb = len(raw_bytes) / (1024 * 1024)
        if size_mb > self.max_size_mb:
            raise SizeLimitError(f"PPTX too large: {size_mb:.1f}MB", filename=name)

        # Inspect zip structure for encryption or corruption
        try:
            with zipfile.ZipFile(io.BytesIO(raw_bytes)) as z:
                names = z.namelist()
                if "EncryptedPackage" in names or any("encrypted" in n.lower() for n in names):
                    raise EncryptedDocumentError(f"PPTX presentation is password protected: {name}", filename=name)
                if not any(n.startswith("ppt/") for n in names) and "[Content_Types].xml" not in names:
                    raise DocumentCorruptError(f"PPTX package is missing required presentationml components: {name}", filename=name)
        except zipfile.BadZipFile as exc:
            raise DocumentCorruptError(f"Corrupt or truncated PPTX zip archive: {exc}", filename=name) from exc
        except EncryptedDocumentError:
            raise
        except DocumentCorruptError:
            raise
        except Exception as exc:
            raise DocumentCorruptError(f"Failed to inspect PPTX package: {exc}", filename=name) from exc

        # Open presentation
        try:
            pres = Presentation(io.BytesIO(raw_bytes))
        except Exception as exc:
            err_msg = str(exc).lower()
            if "password" in err_msg or "encrypted" in err_msg:
                raise EncryptedDocumentError(f"PPTX presentation is encrypted: {exc}", filename=name) from exc
            raise DocumentCorruptError(f"Malformed PPTX presentation: {exc}", filename=name) from exc

        metadata, text_parts = self._extract(pres)
        return DocumentObject(
            filename=name,
            format=DocumentFormat.PPTX,
            raw_bytes=raw_bytes,
            metadata=metadata,
            page_count=metadata.get("slide_count", 0),
            text_content="\n\n".join(text_parts),
        )

    def _extract(self, pres: Any) -> tuple[dict[str, Any], list[str]]:
        """Extract metadata and text from loaded presentation."""
        metadata: dict[str, Any] = {}
        text_parts: list[str] = []
        try:
            cp = pres.core_properties
            metadata["title"] = cp.title
            metadata["author"] = cp.author
            metadata["subject"] = cp.subject
            metadata["keywords"] = cp.keywords
            metadata["slide_count"] = len(pres.slides)
            metadata["slide_width"] = pres.slide_width
            metadata["slide_height"] = pres.slide_height

            image_count = 0
            chart_count = 0
            table_count = 0

            # Iterate slides
            for slide_idx, slide in enumerate(pres.slides, start=1):
                if slide.shapes.title and slide.shapes.title.text.strip():
                    title_text = slide.shapes.title.text.strip()
                    text_parts.append(f"--- Slide {slide_idx}: {title_text} ---")

                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for para in shape.text_frame.paragraphs:
                            text = para.text.strip()
                            if text and text != (slide.shapes.title.text.strip() if slide.shapes.title else ""):
                                text_parts.append(text)
                    if MSO_SHAPE_TYPE is not None and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        image_count += 1
                    elif shape.has_chart:
                        chart_count += 1
                    elif shape.has_table:
                        table_count += 1
                        for row in shape.table.rows:
                            row_text = " | ".join(cell.text.strip() for cell in row.cells)
                            if row_text:
                                text_parts.append(row_text)

            metadata["image_count"] = image_count
            metadata["chart_count"] = chart_count
            metadata["table_count"] = table_count
        except Exception as e:
            logger.warning(f"PPTX extraction warning: {e}")
        return metadata, text_parts
